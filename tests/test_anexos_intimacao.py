from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from unittest.mock import patch
import pytest

from backend.app.servicos import anexos_intimacao as a

IN='IN00000001C'
RECIBO='Protocolo........:123.456\nData Entrada........: 30/10/2026\nGerado via SAEC - Protocolo: IN00000001C - Matrícula/Registro: 1.234;'
OFICIO='''Ofício nº 100100/2026
1. BANCO EXEMPLO - CNPJ 00.360.305/0001-04
Fiduciante(s): PESSOA EXEMPLO - CPF:11144477735
Matrícula(s): 1234
Contrato: 123456789012, firmado em:01/01/2020
Endereço: Rua A, 1, Centro, Morrinhos-GO, CEP 75650082
'''
PROJECAO='''Contrato de financiamento imobiliário Matrícula do Imóvel Matrícula da Garagem
123456789012 1234
Devedores:
PESSOA EXEMPLO - CPF:11144477735
Endereço do imóvel:
Rua A, 1, Centro, Morrinhos-GO, CEP 75650082
Relação de projeção:
30/10/2026 10.545,44
'''


def executar(textos=None, anexos=None, **opcoes):
    textos=textos or dict(recibo=RECIBO,oficio=OFICIO,projecao=PROJECAO)
    anexos=anexos or {p:(p+'.pdf',('%PDF '+p).encode()) for p in a.PAPEIS}
    with patch.object(a, 'ler_anexo', side_effect=lambda corpo,nome,papel,*_: (
        textos[papel],dict(papel=papel,nome=nome,sha256=a.sha256(corpo).hexdigest(),ocr=False))):
        return a.analisar_tres_anexos(anexos,date(2026,10,30),IN,**opcoes)


def test_vincula_tres_fontes_preserva_data_e_valor():
    d=executar()
    assert d['protocoloTri7']=='123456'
    assert d['dataProtocolo']=='30/10/2026'
    assert d['contrato']['data']=='01/01/2020'
    assert d['projecao']['valor']=='10545.44'
    assert len(d['fontes'])==3 and len(d['origens'])>=6
    assert not d['titulo'] and not d['dataTituloAutuacao'] and not d['registroGarantia']


@pytest.mark.parametrize('faltante',a.PAPEIS)
def test_tres_anexos_obrigatorios(faltante):
    arquivos={p:(p+'.pdf',b'%PDF') for p in a.PAPEIS if p!=faltante}
    with pytest.raises(a.Erro,match='obrigatórios'):
        a.analisar_tres_anexos(arquivos,date(2026,10,30),IN)


@pytest.mark.parametrize('papel',a.PAPEIS)
def test_documento_de_outro_in_rejeitado(papel):
    ts=dict(recibo=RECIBO,oficio=OFICIO,projecao=PROJECAO)
    ts[papel]+='\nIN00000002C'
    with pytest.raises(a.Erro,match='outro IN'):
        executar(ts)


def test_recibo_precisa_confirmar_in_e_protocolo_vinculado():
    with pytest.raises(a.Erro,match='diverge'):
        executar(protocolo_tri7='123999')
    with pytest.raises(a.Erro,match='não confirma'):
        executar(dict(recibo=RECIBO.replace(IN,''),oficio=OFICIO,projecao=PROJECAO))


def test_contrato_divergente_nao_e_escolhido_silenciosamente():
    d=executar(dict(recibo=RECIBO,oficio=OFICIO,projecao=PROJECAO.replace('123456789012','999999999999')))
    assert d['contrato']['numero']==''
    assert any(i['campo']=='contrato' for i in d['inconsistencias'])


def test_devedores_divergentes_exigem_correcao():
    d=executar(dict(recibo=RECIBO,oficio=OFICIO,projecao=PROJECAO.replace('PESSOA EXEMPLO','OUTRA PESSOA')))
    assert d['devedores']==[]
    assert any(i['campo']=='devedores' for i in d['inconsistencias'])


def test_valores_divergentes_e_data_futura_nao_sao_escolhidos():
    d=executar(dict(recibo=RECIBO,oficio=OFICIO,projecao=PROJECAO+'30/10/2026 20.000,00\n'))
    assert d['projecao']['valor']=='' and d['projecao']['divergenciaNaData']
    d=executar(dict(recibo=RECIBO,oficio=OFICIO,projecao=PROJECAO.replace('30/10/2026','31/10/2026')))
    assert d['projecao']['valor']==''


def test_excel_xlsx_datas_e_centavos_sem_arredondamento():
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    ws.append(['Data Recebimento','Valor Purga do Débito'])
    ws.append([datetime(2026,10,30),10545.44])
    arquivo=BytesIO();wb.save(arquivo);wb.close()
    texto,fonte=a.ler_anexo(arquivo.getvalue(),'projecao.xlsx','projecao')
    assert '30/10/2026 10545,44' in texto and not fonte['ocr']
    assert a._celula(10.001)=='10,001'
    assert a._celula(123456789012.0)=='123456789012'
    assert a._celula(123456789,'00000000000')=='00123456789'
    assert a._celula(10000,'#,##0.00')=='10000,00'


def test_excel_moeda_inteira_e_numero_de_contrato_sao_preservados():
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active
    ws.append(['Contrato',123456789012])
    ws.append([datetime(2026,10,30),10000]);ws['B2'].number_format='#,##0.00'
    arquivo=BytesIO();wb.save(arquivo);wb.close()
    texto,_=a.ler_anexo(arquivo.getvalue(),'projecao.xlsx','projecao')
    assert 'Contrato 123456789012\n' in texto and '30/10/2026 10000,00' in texto


def test_xml_com_entidades_nao_e_processado():
    from zipfile import ZipFile
    arquivo=BytesIO()
    with ZipFile(arquivo,'w') as z:z.writestr('xl/workbook.xml',b'<!DOCTYPE workbook [<!ENTITY teste "proibido">]><workbook/>')
    with pytest.raises(a.Erro,match='XML'):
        a.ler_anexo(arquivo.getvalue(),'projecao.xlsx','projecao')


@pytest.mark.parametrize('nome,papel,conteudo',[('../origem.pdf','recibo',b'%PDF'),('projecao.xls','projecao',b'<html>fake</html>'),('oficio.xlsx','oficio',b'PK'),('recibo.pdf','recibo',b'')])
def test_rejeita_caminho_formato_inadequado_ou_vazio(nome,papel,conteudo):
    with pytest.raises(a.Erro):a.ler_anexo(conteudo,nome,papel)


def test_pdf_digitalizado_direciona_para_executor_sem_ia():
    with patch.object(a.documentos,'_ler_pdf',side_effect=a.Erro('O PDF parece ser imagem')):
        with pytest.raises(a.Erro,match='executor Windows'):
            a.ler_anexo(b'%PDF teste','recibo.pdf','recibo')


def test_ocr_local_limpa_temporarios_inclusive_em_falha():
    from backend.app.contratos_nucleo import ocr
    arquivos=[]
    def ler(p):
        arquivos.append(p);assert p.exists();return RECIBO
    with patch.object(a.documentos,'_ler_pdf',side_effect=a.Erro('O PDF parece ser imagem')), \
            patch.object(ocr,'disponivel',return_value=True),patch.object(ocr,'texto_de',side_effect=ler):
        _,fonte=a.ler_anexo(b'%PDF teste','recibo.pdf','recibo',True)
    assert fonte['ocr'] and all(not p.exists() for p in arquivos)


def test_pasta_seletor_sem_caminhos_e_somente_arquivos_autorizados(tmp_path):
    pasta=tmp_path/IN;pasta.mkdir();recebidos=pasta/'Recebido';recebidos.mkdir()
    (pasta/'123.456.pdf').write_bytes(b'%PDF recibo')
    (recebidos/'Oficio.pdf').write_bytes(b'%PDF oficio')
    (recebidos/'Projecao.xlsx').write_bytes(b'PK')
    mapa=a.arquivos_da_pasta(IN,pasta=pasta)
    assert len(mapa)==3
    assert {i['sugerido'] for _,i in mapa.values()}==set(a.PAPEIS)
    assert all(set(i)=={'id','nome','tamanho','sugerido'} for _,i in mapa.values())
    with pytest.raises(a.Erro,match='distintos'):
        a.ler_selecao_pasta(IN,dict.fromkeys(a.PAPEIS,next(iter(mapa))),date(2026,10,30),pasta=pasta)
    with pytest.raises(a.Erro,match='não está mais disponível'):
        a.ler_selecao_pasta(IN,dict(zip(a.PAPEIS,('a','b','c'))),date(2026,10,30),pasta=pasta)


def test_pasta_ausente_ou_outro_in(tmp_path):
    with pytest.raises(a.Erro,match='inexistente'):
        a.arquivos_da_pasta(IN,pasta=tmp_path/'ausente')
    with pytest.raises(a.Erro,match='não corresponde'):
        a.arquivos_da_pasta(IN,pasta=tmp_path)


def test_endpoints_novos_exigem_sessao():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from backend.app.rotas.intimacoes import router,preparar_banco
    app=FastAPI();app.include_router(router);app.dependency_overrides[preparar_banco]=lambda:None
    with TestClient(app) as c:
        assert c.post('/api/intimacoes/00000000-0000-0000-0000-000000000001/documentos-evento',json={}).status_code in {401,403}


@pytest.fixture
def cliente_autorizado():
    from fastapi import FastAPI
    from fastapi.testclient import TestClient
    from backend.app.rotas import intimacoes as rotas
    app=FastAPI();app.include_router(rotas.router)
    app.dependency_overrides[rotas.preparar_banco]=lambda:None
    app.dependency_overrides[rotas.proteger_csrf]=lambda:None
    for rota in rotas.router.routes:
        for dep in rota.dependant.dependencies:
            if dep.call.__name__=='verificar':app.dependency_overrides[dep.call]=lambda:'OPERADOR_TESTE'
    with TestClient(app) as c:yield c


def corpo_evento():
    return dict(tipo='oficio',acao='revisao',revisado=True,numeroOficio='INT/2026/053',
                fontes=['1'*64,'2'*64,'3'*64],rascunhoHash='4'*64)


def test_auditoria_de_revisao_nao_reserva_numero_nem_grava_texto(cliente_autorizado):
    from unittest.mock import MagicMock
    from backend.app.rotas import intimacoes as r
    con=MagicMock();con.__enter__.return_value=con
    cur=con.cursor.return_value.__enter__.return_value
    cur.fetchone.side_effect=[{'id':'teste'},{'possivel_duplicidade':False}]
    with patch.object(r,'conectar',return_value=con),patch.object(r,'registrar_auditoria_cursor') as audit:
        resposta=cliente_autorizado.post('/api/intimacoes/00000000-0000-0000-0000-000000000001/documentos-evento',json=corpo_evento())
    assert resposta.status_code==200 and not resposta.json()['numeroReservado']
    detalhes=audit.call_args.args[-1]
    assert detalhes['rascunhoHash']=='4'*64
    assert not detalhes['usoNaTri7Confirmado'] and 'texto' not in detalhes


def test_duplicidade_de_copia_em_outro_in_exige_conferencia(cliente_autorizado):
    from unittest.mock import MagicMock
    from backend.app.rotas import intimacoes as r
    con=MagicMock();con.__enter__.return_value=con
    con.cursor.return_value.__enter__.return_value.fetchone.side_effect=[{'id':'teste'},{'possivel_duplicidade':True}]
    with patch.object(r,'conectar',return_value=con),patch.object(r,'registrar_auditoria_cursor') as audit:
        resposta=cliente_autorizado.post('/api/intimacoes/00000000-0000-0000-0000-000000000001/documentos-evento',json=corpo_evento())
    assert resposta.status_code==409
    audit.assert_not_called()


def test_revisao_obrigatoria_e_tamanho_formato_de_json(cliente_autorizado):
    url='/api/intimacoes/00000000-0000-0000-0000-000000000001/documentos-evento'
    corpo=corpo_evento();corpo['revisado']=False
    assert cliente_autorizado.post(url,json=corpo).status_code==422
    assert cliente_autorizado.post(url,content='x'*2049,headers={'Content-Type':'application/json'}).status_code==413
    assert cliente_autorizado.post(url,content='{}',headers={'Content-Type':'text/plain'}).status_code==415


def test_recibo_e_obrigatorio_tambem_na_api(cliente_autorizado):
    url='/api/intimacoes/00000000-0000-0000-0000-000000000001/preparar-documentos'
    resposta=cliente_autorizado.post(url,files={'oficio':('oficio.pdf',b'%PDF'), 'projecao':('projecao.pdf',b'%PDF')})
    assert resposta.status_code==422

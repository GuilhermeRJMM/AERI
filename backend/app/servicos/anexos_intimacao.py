"""Leitura dos três anexos do módulo existente; sem IA, envio ou gravação de originais."""
from datetime import date, datetime
from decimal import Decimal
from hashlib import sha256
from io import BytesIO, StringIO
from itertools import chain
from pathlib import Path
import re
import tempfile
from zipfile import ZipFile, BadZipFile

from backend.app.servicos import documentos_intimacao as documentos

Erro = documentos.ErroDocumentoIntimacao
PAPEIS = ('recibo', 'oficio', 'projecao')
LIMITE = documentos.LIMITE_PDF_INTIMACAO
VERSAO = 'tres-anexos-20261030'


def _matricula(texto):
    recibo = re.search(r'Matr[ií]cula(?:/Registro|\(s\)|s)?\s*:\s*([\d.]+)', texto, re.I)
    return recibo.group(1).strip('.') if recibo else _cabecalho_projecao(texto).get('matricula', '')


def _cabecalho_projecao(texto):
    m = re.search(r'Contrato de financiamento imobili[aá]rio\s+Matr[ií]cula do Im[oó]vel[^\n]*\n\s*([\d.\-/]+)\s+([\d.]+)', texto, re.I)
    return {'contrato': m.group(1), 'matricula': m.group(2)} if m else {}


def _contrato(texto):
    return _dados_contrato(texto)['numero'] or _cabecalho_projecao(texto).get('contrato', '')


def _dados_contrato(texto):
    contrato = documentos._extrair_contrato(texto)
    if contrato['numero']:
        return contrato
    m = re.search(r'Contrato\s*(?:n[ºo°.]*\s*)?:\s*(\d[\d./-]*)', texto, re.I)
    data = re.search(r'(?:Data de Assinatura|firmad[oa]s?\s+em)\s*:?\s*(\d{2}/\d{2}/\d{4})', texto[m.end():m.end()+180], re.I) if m else None
    return dict(numero=m.group(1) if m else '', data=data.group(1) if data else '')


def _grupos_cotas(texto):
    m = re.search(r'\bGrupo(?:\(\s*s\s*\)|s)?\s*[:/]', texto, re.I)
    return {documentos._digitos(n) for n in re.findall(r'\b\d{4,8}\s*/\s*\d{3,6}\b', texto[max(0,m.start()-250):m.start()+700])} if m else set()


def _credor(texto):
    # Qualificação narrativa: nome da pessoa jurídica antes da sede e CNPJ.
    for nome in re.finditer(r"(?m)^\s*([A-ZÀ-ÖØ-Þ][A-ZÀ-ÖØ-Þ \t.&/-]{4,160}?\b(?:LTDA|S/?A|S\.A\.|EIRELI))\s*[,;]", texto):
        cnpj = re.search(r'CNPJ\s*(?:sob\s*)?(?:n[.ºo°]*\s*)?[:.]?\s*([\d./-]{14,20})', texto[nome.end():nome.end()+500], re.I)
        if cnpj and len(documentos._digitos(cnpj.group(1))) == 14:
            return dict(nome=nome.group(1).strip(), cnpj=documentos._digitos(cnpj.group(1)))
    return documentos._extrair_credor(texto)


def _devedores(texto):
    ps = documentos._extrair_devedores(texto)
    bloco = re.search(r'Devedores?\s*:\s*(.*?)(?=Endere[cç]o do im[oó]vel|\Z)', texto, re.I | re.S)
    if bloco:
        ps = ps or documentos._extrair_devedores('Fiduciante(s): ' + bloco.group(1))
    inicio = re.search(r'Fiduciante(?:\(s\)|s)?\s*:', texto, re.I)
    if not inicio:
        return ps
    trecho = re.split(r'CL[AÁ]USULA|Dados do Im[oó]vel|Matr[ií]cula\s*:', texto[inicio.end():inicio.end()+5000], maxsplit=1, flags=re.I)[0]
    padrao = r"([A-ZÀ-ÖØ-Þ][A-ZÀ-ÖØ-Þ .'-]{3,100}?)\s*,\s*(?i:brasileir[oa]|casad[oa]|solteir[oa]|divorciad[oa]|portador[ae]?)\b.{0,700}?\bCPF\s*(?i:sob\s*)?(?i:n[ºo°.]*\s*)?[:.]?\s*([\d. -]{11,18})"
    for m in re.finditer(padrao, trecho, re.S):
        pessoa = dict(nome=documentos._nome_arquivo(m.group(1)), cpf=documentos._digitos(m.group(2)))
        if len(pessoa['cpf']) != 11:
            continue
        if any(p.get('cpf') == pessoa['cpf'] and documentos._nome_para_comparacao(p['nome']) != documentos._nome_para_comparacao(pessoa['nome']) for p in ps):
            raise Erro('O ofício apresenta nomes diferentes para o mesmo CPF. Confira o original.')
        if pessoa not in ps:
            ps.append(pessoa)
    return ps


def _texto_corrompido(texto):
    ruins = sum(c in '\ufffd\u25a1' or ord(c) < 32 and c not in '\n\r\t' for c in texto)
    return '\ufffd' in texto or ruins >= 5 and ruins / max(len(texto), 1) > .02


def _celula(valor, formato=''):
    if valor is None:
        return ''
    if isinstance(valor, (date, datetime)):
        return valor.strftime('%d/%m/%Y')
    if isinstance(valor, (int, float, Decimal)) and not isinstance(valor, bool):
        # Só muda a representação decimal; não arredonda nem calcula valores.
        numero = Decimal(str(valor))
        if not numero.is_finite():
            raise Erro('A planilha apresenta um número inválido.')
        if numero == numero.to_integral_value() and not re.search(r'[0#][.,][0#]{2}', formato):
            if re.fullmatch('0{4,20}', formato):
                return str(int(numero)).zfill(len(formato))
            return str(int(numero))
        return format(numero, '.2f' if numero.as_tuple().exponent >= -2 else 'f').replace('.', ',')
    return str(valor)


def _ler_excel(conteudo, extensao):
    linhas = []
    celulas = 0
    try:
        if extensao == '.xlsx':
            from openpyxl import load_workbook
            with ZipFile(BytesIO(conteudo)) as pacote:
                itens = pacote.infolist()
                if len(itens) > 2000 or sum(i.file_size for i in itens) > 20_000_000:
                    raise Erro('A planilha excede o limite de leitura descompactada.')
                for item in itens:
                    if item.filename.endswith('.xml'):
                        xml = pacote.read(item).upper()
                        if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:
                            raise Erro('A estrutura XML da planilha não é permitida.')
            livro = load_workbook(BytesIO(conteudo), read_only=True, data_only=True, keep_links=False)
            try:
                if len(livro.worksheets) > 20:
                    raise Erro('A planilha contém mais de 20 abas.')
                for folha in livro.worksheets:
                    if (folha.max_row or 0) > 5000 or (folha.max_column or 0) > 100:
                        raise Erro('A planilha excede 5.000 linhas ou 100 colunas por aba.')
                    for indice, linha in enumerate(folha.iter_rows()):
                        celulas += len(linha)
                        if indice >= 5000 or len(linha) > 100 or celulas > 100_000:
                            raise Erro('A planilha excede o limite de células da leitura.')
                        linhas.append(' '.join(_celula(c.value, getattr(c, 'number_format', '')) for c in linha))
            finally:
                livro.close()
        elif extensao == '.xls' and conteudo.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):
            import xlrd
            livro = xlrd.open_workbook(file_contents=conteudo, logfile=StringIO(), on_demand=True, formatting_info=True)
            try:
                if livro.nsheets > 20:
                    raise Erro('A planilha contém mais de 20 abas.')
                for folha in livro.sheets():
                    if folha.nrows > 5000 or folha.ncols > 100:
                        raise Erro('A planilha excede o limite de linhas ou colunas.')
                    for linha in folha.get_rows():
                        celulas += len(linha)
                        if celulas > 100_000:
                            raise Erro('A planilha excede o limite de células da leitura.')
                        linhas.append(' '.join(_celula(xlrd.xldate_as_datetime(c.value, livro.datemode)
                            if c.ctype == xlrd.XL_CELL_DATE else c.value,
                            livro.format_map[livro.xf_list[c.xf_index].format_key].format_str
                            if livro.formatting_info else '') for c in linha))
            finally:
                livro.release_resources()
        else:
            raise Erro('Use uma planilha Excel XLS/XLSX válida ou um PDF da projeção.')
    except ImportError:
        raise Erro('A leitura Excel não está instalada neste ambiente; use o executor configurado ou PDF.') from None
    except Erro:
        raise
    except Exception:
        raise Erro('Não foi possível ler a planilha. Confira o formato e os valores calculados salvos no Excel.') from None
    texto = '\n'.join(linhas)
    if len(texto) > 1_000_000:
        raise Erro('O conteúdo da planilha excede o limite de leitura.')
    return texto


def ler_anexo(conteudo, nome, papel, permitir_ocr=False):
    nome = str(nome or '')
    if not nome or Path(nome).name != nome or '/' in nome or '\\' in nome:
        raise Erro('Nome de arquivo inválido.')
    if not conteudo or len(conteudo) > LIMITE:
        raise Erro('Cada anexo deve ser legível e ter até 8 MB.')
    extensao = Path(nome).suffix.lower()
    ocr_usado = False
    if extensao == '.pdf':
        try:
            texto = documentos._ler_pdf(conteudo, papel)
            if _texto_corrompido(texto):
                raise Erro('O PDF parece ser imagem ou tem codificação de texto corrompida.')
        except Erro as exc:
            if 'parece ser imagem' not in str(exc):
                raise
            if not permitir_ocr:
                raise Erro('PDF digitalizado ou com texto corrompido: use a leitura da pasta pelo executor Windows para realizar OCR local.') from None
            from backend.app.contratos_nucleo import ocr
            if not ocr.disponivel():
                raise Erro('OCR local indisponível no executor. Instale/configure o motor ou forneça um PDF pesquisável.')
            # Os únicos arquivos temporários são removidos inclusive em falha.
            with tempfile.TemporaryDirectory(prefix='aeri-intimacao-') as pasta:
                arquivo = Path(pasta) / 'anexo.pdf'
                arquivo.write_bytes(conteudo)
                texto = ocr.texto_de(arquivo, linhas_espaciais=True) if papel == 'projecao' else ocr.texto_de(arquivo)
            ocr_usado = True
    elif papel == 'projecao' and extensao in {'.xlsx', '.xls'}:
        texto = _ler_excel(conteudo, extensao)
    else:
        raise Erro('Recibo e ofício devem ser PDFs; a projeção pode ser PDF, XLS ou XLSX.')
    if len(texto.strip()) < 30 or _texto_corrompido(texto):
        raise Erro('O anexo está ilegível ou contém caracteres corrompidos. Revise o arquivo de origem.')
    return texto, dict(papel=papel, nome=nome, sha256=sha256(conteudo).hexdigest(), ocr=ocr_usado)


def _enderecos_do_oficio(texto):
    """Complemento rural/multilinha restrito à qualificação dos fiduciantes."""
    encontrados = documentos._extrair_enderecos(texto)
    inicio = re.search(r'Fiduciante(?:\(s\)|s)?\s*:', texto, re.I)
    bloco = texto[inicio.end():] if inicio else ''
    bloco = re.split(r'\n\s*(?:3\.1|4\.)\s', bloco, maxsplit=1)[0]
    linhas = bloco.splitlines()
    inicio_endereco = r'(?:Rua|Avenida|Av\.|Alameda|Praça|Estrada|Rodovia|Fazenda|S[ií]tio|Ch[aá]cara|Gleba)\b'
    for i, linha in enumerate(linhas):
        m = re.match(r'^\s*(?:\d+[.)]\s*)?('+inicio_endereco+r'.+)', linha, re.I)
        if not m:
            continue
        partes = [m[1].strip()]
        # Não continuar após município/UF já presente; evita anexar outra pessoa.
        for proxima in linhas[i+1:i+5]:
            if re.search(r'\b[A-ZÀ-Ü][A-ZÀ-Üa-zà-ü ]+\s*[-/]\s*[A-Z]{2}\b|Munic[ií]pio de Morrinhos', ' '.join(partes)):
                break
            trecho = proxima.strip()
            if not trecho or re.match(r'^(?:\d+[.)]|'+inicio_endereco+r'|CPF|CNPJ|Nome|Contrato|Endere[cç]o|Registro)\b', trecho, re.I):
                break
            partes.append(trecho)
        valor = documentos._corrigir_indicador_numero(re.sub(r'\s+', ' ', ' '.join(partes))).strip(' ,;.')
        # Substitui a linha parcial, sem duplicar o mesmo endereço.
        encontrados = [e for e in encontrados if e != partes[0].strip(' ,;.')]
        if valor not in encontrados:
            encontrados.append(valor)
    return encontrados


def analisar_tres_anexos(anexos, referencia, protocolo_in, protocolo_tri7='', permitir_ocr=False):
    if not re.fullmatch(r'IN\d{8}C', protocolo_in):
        raise Erro('Selecione um IN válido cadastrado no AERI.')
    if set(anexos) != set(PAPEIS):
        raise Erro('Recibo do Protocolo, Ofício do Credor e Planilha de Projeção são obrigatórios.')
    textos, fontes = {}, []
    for papel in PAPEIS:
        nome, conteudo = anexos[papel]
        textos[papel], fonte = ler_anexo(conteudo, nome, papel, permitir_ocr)
        fontes.append(fonte)
        ins = {re.sub(r'\s', '', s).upper() for s in re.findall(r'\bIN\s*\d{8}\s*C\b', textos[papel], re.I)}
        if ins and ins != {protocolo_in}:
            raise Erro('Um dos documentos pertence a outro IN. Selecione os três arquivos do procedimento correto.')
    if not re.search(rf'\b{protocolo_in}\b', textos['recibo'], re.I):
        raise Erro('O Recibo do Protocolo não confirma o IN selecionado.')
    recibo = textos['recibo']
    numero = re.search(r'\bProtocolo[.\s:]{0,40}(\d{1,3}(?:\.\d{3})+|\d{3,10})(?!\d)', recibo, re.I)
    entrada = re.search(r'\b(?:Data\s+Entrada|Data\s+do\s+protocolo)[.\s:]{0,40}(\d{2}/\d{2}/\d{4})', recibo, re.I)
    if not entrada:
        entrada = re.search(r'\bProtocolo[^\n]{0,60}\bde\s+(\d{2}/\d{2}/\d{4})', recibo, re.I)
    if not numero or not entrada:
        raise Erro('Recibo inadequado: número e data de entrada do protocolo não foram identificados.')
    protocolo = documentos._digitos(numero.group(1))
    protocolo_divergente = bool(protocolo_tri7 and documentos._digitos(protocolo_tri7) != protocolo)
    try:
        datetime.strptime(entrada.group(1), '%d/%m/%Y')
    except ValueError:
        raise Erro('A data de entrada no recibo é inválida.') from None
    dados = documentos.analisar_pdfs_intimacao(b'', b'', referencia, protocolo_in, protocolo,
                                              textos=(textos['oficio'], textos['projecao']))
    dados.update(fluxo=VERSAO, dataProtocolo=entrada.group(1), fontes=fontes,
                 inconsistencias=[], origens={}, registroGarantia='', titulo='', dataTituloAutuacao='')
    dados['credor'] = _credor(textos['oficio'])
    dados['devedores'] = _devedores(textos['oficio'])
    dados['enderecosDoOficio'] = _enderecos_do_oficio(textos['oficio'])
    dados['matricula'] = _matricula(textos['oficio']) or _matricula(recibo)
    dados['contrato'] = _dados_contrato(textos['oficio'])
    registros = set(re.findall(r'Registro(?:\(s\)|s)?\s*:\s*(R[.\s-]*\d+)', textos['oficio'], re.I))
    dados['registroIndicado'] = next(iter(registros)) if len(registros) == 1 else ''
    outorga = re.search(r'OUTORGA.{0,800}?Cl[áa]usula\s+([IVXLCDM]+|\d+)\b', textos['oficio'], re.I|re.S)
    dados['clausula'] = f'Cláusula {outorga[1]}' if outorga else ''
    dados_imovel = re.search(r'Dados do Im[oó]vel\s+Endere[cç]o\s*:\s*(.*?)(?=\s*Matr[ií]cula\s*:)', textos['oficio'], re.I|re.S)
    if not dados['enderecoImovel'] and dados_imovel and len(dados_imovel[1]) < 1500:
        dados['enderecoImovel'] = re.sub(r'\s+', ' ', dados_imovel[1]).strip()
    if protocolo_divergente:
        fontes_protocolo = {'recibo': protocolo, 'cadastroAeri': str(protocolo_tri7)}
        dados['inconsistencias'].append(dict(campo='protocoloTri7', fontes=fontes_protocolo))
        dados['avisos'].append(f'O recibo deste IN informa protocolo {protocolo}; o cadastro antigo do AERI contém {protocolo_tri7}. Confira a divergência antes de aprovar. O rascunho usa o recibo e o cadastro NÃO foi alterado.')
    if len({f['sha256'] for f in fontes}) != 3:
        raise Erro('Os três anexos precisam ser documentos distintos.')
    if not _devedores(textos['oficio']) or not dados['credor']['nome']:
        raise Erro('Ofício inadequado: credor e devedores não puderam ser identificados.')
    if not dados['projecao']['datasDisponiveis']:
        raise Erro('A projeção não apresenta linhas legíveis de data e valor. Salve os cálculos no Excel ou forneça o PDF pesquisável.')
    dados['origens'].update(protocoloTri7={'recibo': protocolo, 'cadastroAeri': str(protocolo_tri7 or '')}, dataProtocolo={'recibo': entrada.group(1)},
        credor={'oficio': dados['credor']}, devedores={'oficio': dados['devedores']},
        projecao={'projecao': dados['projecao']}, enderecoImovel={'oficio/projecao': dados['enderecoImovel']})

    def comparar(campo, extrator, limpar):
        valores = {papel: extrator(texto) for papel, texto in textos.items()}
        valores = {p: v for p, v in valores.items() if v}
        dados['origens'][campo] = valores
        if len({limpar(v) for v in valores.values()}) > 1:
            dados['inconsistencias'].append(dict(campo=campo, fontes=valores))
            if campo == 'matricula':
                dados['matricula'] = ''
            elif campo == 'contrato':
                dados['contrato']['numero'] = ''
            elif campo == 'enderecoImovel':
                dados['enderecoImovel'] = ''
        return valores

    comparar('matricula', _matricula, documentos._digitos)
    comparar('contrato', _contrato, documentos._digitos)
    comparar('enderecoImovel', documentos._extrair_endereco_imovel,
             lambda v: re.sub(r'\W', '', documentos._nome_para_comparacao(v)))
    datas = comparar('dataContrato', lambda t: _dados_contrato(t)['data'], str)
    if len(set(datas.values())) > 1:
        dados['contrato']['data'] = ''
    pessoas = {p: _devedores(t) for p, t in textos.items() if p != 'recibo'}
    conjuntos = [{(documentos._nome_para_comparacao(x['nome']), x['cpf']) for x in ps} for ps in pessoas.values() if ps]
    if len(conjuntos) == 2 and conjuntos[0] != conjuntos[1]:
        dados['inconsistencias'].append(dict(campo='devedores', fontes=pessoas))
        dados['devedores'] = []
    # Sem o IN nos documentos do credor, exigimos vínculo cruzado verificável.
    matricula_recibo = _matricula(recibo)
    for papel in ('oficio', 'projecao'):
        contem_in = protocolo_in in textos[papel].upper()
        matricula = _matricula(textos[papel])
        contrato = _contrato(textos[papel])
        mesma_matricula = bool(matricula_recibo and matricula and documentos._digitos(matricula_recibo) == documentos._digitos(matricula))
        if matricula_recibo and matricula and not mesma_matricula:
            raise Erro('A matrícula de um dos anexos diverge da matrícula do recibo. Confira os arquivos do IN.')
        oficio_vinculado = bool(matricula_recibo and documentos._digitos(matricula_recibo) == documentos._digitos(_matricula(textos['oficio'])))
        contrato_oficio = documentos._digitos(_contrato(textos['oficio']))
        mesmos_grupos = bool(papel == 'projecao' and oficio_vinculado and contrato_oficio and
                             contrato_oficio in _grupos_cotas(textos[papel]) and
                             contrato_oficio in _grupos_cotas(textos['oficio']))
        mesmo_contrato = bool(papel == 'projecao' and oficio_vinculado and contrato and contrato_oficio and documentos._digitos(contrato) == contrato_oficio)
        if not contem_in and not (mesma_matricula or mesmos_grupos or mesmo_contrato):
            raise Erro('Não foi possível vincular todos os anexos ao IN pela matrícula do recibo. Confira os arquivos.')
        if not (matricula or contrato or mesmos_grupos):
            raise Erro('Ofício ou projeção sem identificação mínima do contrato ou da matrícula.')
    dados['avisos'].append('Título, data do título na Autuação, registro da garantia, selo e cláusula contratual exigem conferência nos documentos; não foram presumidos.')
    if any(f['ocr'] for f in fontes):
        dados['avisos'].append('Há dados obtidos por OCR local: confira nomes, documentos, datas e valores no original.')
    if 'CONSORCIO' in documentos._nome_para_comparacao(dados['credor']['nome']).upper():
        dados['avisos'].append('Credor de consórcio: confira a aplicabilidade dos modelos e das cláusulas. Não trate o instrumento como financiamento residencial por presunção.')
    return dados


def arquivos_da_pasta(protocolo, pasta=None):
    from ferramentas.abrir_pasta_intimacao import localizar_pasta_existente, RAIZES_BUSCA
    raiz = Path(pasta) if pasta is not None else localizar_pasta_existente(protocolo)
    if not re.fullmatch(r'IN\d{8}C', protocolo):
        raise Erro('IN inválido.')
    if raiz is None or not raiz.is_dir():
        raise Erro('Pasta deste IN inexistente ou inacessível ao executor.')
    raiz = raiz.resolve()
    if pasta is None and not any(raiz.is_relative_to(Path(r).resolve()) for r in RAIZES_BUSCA):
        raise Erro('A pasta está fora das raízes autorizadas.')
    if raiz.name.upper() != protocolo:
        raise Erro('A pasta localizada não corresponde ao IN selecionado.')
    mapa = {}
    # Apenas a pasta do IN e suas subpastas imediatas, nunca a rede inteira.
    for arquivo in chain(raiz.iterdir(), chain.from_iterable(sub.iterdir() for sub in raiz.iterdir()
                         if sub.is_dir() and not sub.is_symlink())):
        if not arquivo.is_file() or arquivo.is_symlink() or not arquivo.resolve().is_relative_to(raiz):
            continue
        if arquivo.suffix.lower() not in {'.pdf', '.xls', '.xlsx'}:
            continue
        relativo = arquivo.relative_to(raiz).as_posix()
        chave = sha256(relativo.encode('utf-8')).hexdigest()
        nome = arquivo.stem.upper()
        sugerido = ('recibo' if re.fullmatch(r'\d{1,3}\.\d{3}', nome) or 'RECIBO' in nome
                    else 'projecao' if 'PROJE' in nome else 'oficio' if 'OFICIO' in nome or 'OFÍCIO' in nome else '')
        if 'VALIDA' in nome or 'CONTRA' in nome:
            sugerido = ''
        mapa[chave] = (arquivo, dict(id=chave, nome=arquivo.name, tamanho=arquivo.stat().st_size, sugerido=sugerido))
        if len(mapa) > 200:
            raise Erro('Há muitos arquivos nesta pasta. Organize os anexos antes da leitura.')
    return mapa


def ler_selecao_pasta(protocolo, selecao, referencia, protocolo_tri7='', pasta=None):
    mapa = arquivos_da_pasta(protocolo, pasta)
    if set(selecao) != set(PAPEIS) or len(set(selecao.values())) != 3:
        raise Erro('Selecione três arquivos distintos: recibo, ofício e projeção.')
    anexos = {}
    for papel, chave in selecao.items():
        if chave not in mapa:
            raise Erro('Um arquivo não está mais disponível na pasta autorizada. Localize novamente.')
        arquivo = mapa[chave][0]
        raiz = arquivo.parent if arquivo.parent.name.upper() == protocolo else arquivo.parent.parent
        if arquivo.is_symlink() or not arquivo.resolve().is_relative_to(raiz.resolve()):
            raise Erro('O arquivo saiu da pasta autorizada. Localize os anexos novamente.')
        if arquivo.stat().st_size > LIMITE:
            raise Erro('Um anexo ultrapassa 8 MB.')
        with arquivo.open('rb') as stream:
            conteudo = stream.read(LIMITE + 1)
        if len(conteudo) > LIMITE:
            raise Erro('Um anexo ultrapassa 8 MB.')
        anexos[papel] = (arquivo.name, conteudo)
    return analisar_tres_anexos(anexos, referencia, protocolo, protocolo_tri7, permitir_ocr=True)
